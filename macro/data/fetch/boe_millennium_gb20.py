"""Great Britain around the crisis of 1825: monthly share prices, Bank Rate, the market rate of discount and
wholesale prices (1822-1827), and annual output by sector and by industry (1815-1830), from the Bank of
England's "A millennium of macroeconomic data" (Chapter 20).

    python3 data/fetch/boe_millennium_gb20.py           # download the workbook (27 MB) and write the snapshots
    python3 data/fetch/boe_millennium_gb20.py FILE      # use a copy of the workbook already downloaded

Monthly, 1822-1827 (data/snapshots/boe_millennium_gb_1822_1827_monthly.csv):
  sheet "M13. Mthly share prices 1709+": the Gayer, Rostow and Schwartz (1953) share price indices of 63
  companies, inclusive of mines and excluding mines (June 1840 = 100), and Hayek's unweighted index (1841 = 100);
  sheet "M9. Mthly short-term rates": Bank Rate (end month) and the rate of discount for three-month prime bills
  (Overend Gurney, monthly average of weekly data, from 1824);
  sheet "M6. Mthly prices and wages": the Gayer-Rostow-Schwartz wholesale price index (1821-5 = 100).
Annual, 1815-1830 (data/snapshots/boe_millennium_gb_1815_1830.csv):
  sheet "A7. GB GDP(O) 1700-1870": real GDP of Great Britain by sector (1700 = 100), nominal GDP (£ m) and the
  GDP deflator (1700 = 100), Broadberry et al. (2015);
  sheet "A4. Ind Production 1270-1870": output of eight key industries (Broadberry et al. 2015, Table 4.02).
The script checks every heading it relies on and fails if one has moved.
"""
from __future__ import annotations

import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
SNAP = HERE.parent / "snapshots"
OUT_M = SNAP / "boe_millennium_gb_1822_1827_monthly.csv"
OUT_A = SNAP / "boe_millennium_gb_1815_1830.csv"
URL = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/"
       "a-millennium-of-macroeconomic-data-for-the-uk.xlsx")
UA = "Mozilla/5.0 (X11; Linux x86_64) finance-study data fetch"
M_FIRST, M_LAST = 1822, 1827
A_FIRST, A_LAST = 1815, 1830
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
M13, M9, M6 = "M13. Mthly share prices 1709+ ", "M9. Mthly short-term rates", "M6. Mthly prices and wages"
A7, A4 = "A7. GB GDP(O) 1700-1870", "A4. Ind Production 1270-1870"
# monthly series: (sheet, column, [(row, expected heading)], name in the snapshot)
MONTHLY = [
    (M13, 3, [(2, "Gayer, Rostow and Schwarz (1953) index, 1811-1850"),
              (3, "63 companies, inclusive of mines, June 1840=100, weighted by numbers of shares/paid up capital")],
     "shares_incl_mines"),
    (M13, 4, [(2, "Gayer, Rostow and Schwarz (1953) index, 1811-1850"),
              (3, "63 companies, excluding mines, June 1840=100, weighted by numbers of shares/paid up capital")],
     "shares_excl_mines"),
    (M13, 5, [(2, "Hayek's index of Share prices, 1820-1868"), (3, "Unweighted index, 1841=100")], "shares_hayek"),
    (M9, 2, [(5, "D1. Official rates"), (6, "End month")], "bank_rate"),
    (M9, 6, [(4, "Rate of  discount for three month prime or first class bills 1824-1939"),
             (6, "Monthly average of weekly data")], "prime_bills_rate"),
    (M6, 4, [(4, "Domestic and imported commodities 1790-1850"), (5, "Gayer et al (1953)"), (6, "1821-5=100")],
     "wholesale_1821_25"),
]
# annual series: (sheet, column, [(row, expected heading)], name in the snapshot)
ANNUAL = [
    (A7, 1, [(6, "Real GDP (1700 =100)"), (7, "Agriculture")], "real_agriculture"),
    (A7, 2, [(7, "Industry")], "real_industry"),
    (A7, 3, [(7, "Services")], "real_services"),
    (A7, 4, [(7, "GDP")], "real_gdp"),
    (A7, 6, [(6, "Nominal GDP (mln £)")], "nominal_gdp_gbp_m"),
    (A7, 11, [(7, "GDP Deflator")], "gdp_deflator"),
] + [(A4, c, ([(6, "Table 4.02:  Output of key industries")] if c == 1 else []) + [(7, h)], "ind_" + n)
       for c, h, n in [
    (1, "Tin", "tin"), (2, "Iron", "iron"), (3, "Coal", "coal"), (4, "Wool/Textiles", "textiles"),
    (5, "Leather", "leather"), (6, "Foodstuffs", "foodstuffs"), (7, "Construction", "construction"),
    (8, "Printed books", "books")]]


def year_of(cell) -> int | None:
    if isinstance(cell, (int, float)):
        return int(cell)
    if isinstance(cell, str) and cell[:4].isdigit():
        return int(cell[:4])
    return None


def check(rows: list, sheet: str, col: int, heads: list[tuple[int, str]]) -> None:
    for row, head in heads:
        got = str(rows[row][col]).strip()
        if got != head.strip():
            raise ValueError(f"{sheet}: column {col} row {row} is {rows[row][col]!r}, expected {head!r}")


def main(path: Path) -> None:
    if not path.exists():
        req = urllib.request.Request(URL, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=300) as r:
            path.write_bytes(r.read())
    print(f"{path.name}: sha256 {hashlib.sha256(path.read_bytes()).hexdigest()}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list] = {}

    monthly: dict[str, dict[tuple[int, int], float]] = {}
    for sheet, col, heads, name in MONTHLY:
        rows = sheets.setdefault(sheet, list(wb[sheet].iter_rows(values_only=True)))
        check(rows, sheet, col, heads)
        s = {}
        for r in rows[7:]:
            y = year_of(r[0])
            if y is not None and M_FIRST <= y <= M_LAST and r[1] in MONTHS and isinstance(r[col], (int, float)):
                s[(y, MONTHS.index(r[1]) + 1)] = float(r[col])
        monthly[name] = s
    keys = [(y, m) for y in range(M_FIRST, M_LAST + 1) for m in range(1, 13)]
    with OUT_M.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "month"] + list(monthly))
        for k in keys:
            w.writerow(list(k) + ["" if k not in monthly[n] else round(monthly[n][k], 6) for n in monthly])
    print(f"{len(keys)} months -> {OUT_M}")
    for n, s in monthly.items():
        print(f"  {n}: {min(s)}-{max(s)}, {len(s)} values")

    annual: dict[str, dict[int, float]] = {}
    for sheet, col, heads, name in ANNUAL:
        rows = sheets.setdefault(sheet, list(wb[sheet].iter_rows(values_only=True)))
        check(rows, sheet, col, heads)
        s = {}
        for r in rows[8:]:
            y = year_of(r[0])
            if y is not None and A_FIRST <= y <= A_LAST and isinstance(r[col], (int, float)):
                s.setdefault(y, float(r[col]))
        annual[name] = s
    years = range(A_FIRST, A_LAST + 1)
    with OUT_A.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year"] + list(annual))
        for y in years:
            w.writerow([y] + ["" if y not in annual[n] else round(annual[n][y], 6) for n in annual])
    print(f"{len(years)} years -> {OUT_A}")
    for n, s in annual.items():
        print(f"  {n}: {min(s)}-{max(s)}, {len(s)} values")


if __name__ == "__main__":
    default = Path.home() / ".cache" / "finance-study" / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    p.parent.mkdir(parents=True, exist_ok=True)
    main(p)
