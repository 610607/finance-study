"""British war finance, 1913-1950: central government expenditure and receipts, the national debt and consumer
prices, from the Bank of England's "A millennium of macroeconomic data" (Chapter 34).

    python3 data/fetch/boe_millennium_gb34.py           # download the workbook (27 MB) and write the snapshot
    python3 data/fetch/boe_millennium_gb34.py FILE      # use a copy of the workbook already downloaded

Annual, 1913-1950 (data/snapshots/boe_millennium_gb_1913_1950.csv), one row per calendar year y:
  sheet "A27. Central govt borrowing ": total expenditure and total receipts of the financial year that begins in
  April of y and ends on 31 March of y+1 (fy '1913/14' in the row of 1913), GBP million (Mitchell 1988 to 1945,
  the ONS from 1946); and the sheet's calendar-year totals of central government expenditure and income as a per
  cent of nominal GDP (the calendar year is in column 43; it sits one row below the financial year that ends in it);
  sheet "A29. The National Debt": the spliced measure of consolidated UK debt, approximate end-of-calendar-year
  estimate, GBP million, and as a per cent of nominal GDP at market prices;
  sheet "A47. Wages and prices": the preferred consumer price index of the UK (2015 = 100), the column of Chapter 27's
  table of 1913-1931.
The script checks every heading it relies on and fails if one has moved. The CSV is written with LF line ends.
"""
from __future__ import annotations

import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from boe_millennium_gb27 import URL, UA, check, number  # noqa: E402

SNAP = HERE.parent / "snapshots"
OUT = SNAP / "boe_millennium_gb_1913_1950.csv"
FIRST, LAST = 1913, 1950
A27, A29, A47 = "A27. Central govt borrowing ", "A29. The National Debt", "A47. Wages and prices"
HEADER = ["year", "fy", "fy_expenditure_gbp_mn", "fy_receipts_gbp_mn", "cy_expenditure_pct_gdp",
          "cy_receipts_pct_gdp", "debt_gbp_mn", "debt_pct_gdp", "cpi_2015_100"]


def write(path: Path, header: list[str], rows: list[list]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def main(path: Path) -> None:
    if not path.exists():
        req = urllib.request.Request(URL, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=300) as r:
            path.write_bytes(r.read())
    print(f"{path.name}: sha256 {hashlib.sha256(path.read_bytes()).hexdigest()}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    years = list(range(FIRST, LAST + 1))
    table: dict[int, dict[str, object]] = {y: {} for y in years}

    # A27: the financial years (columns 0, 2, 13) and the calendar years (columns 43, 46, 49)
    rows = list(wb[A27].iter_rows(values_only=True))
    check(rows, A27, 2, [(0, "A27. Central Government Revenue and Expenditure"),
                         (1, "Source: Mitchell (1988) pages 575-593 for 1688-1945, ONS from 1946 onwards"),
                         (6, "Financial Year Totals"), (7, "Expenditure"), (8, "Total expenditure")])
    check(rows, A27, 13, [(7, "Receipts"), (8, "Total receipts")])
    check(rows, A27, 44, [(6, "Calendar Year Totals")])
    check(rows, A27, 45, [(8, "Expenditure"), (9, "Central govt. expenditure")])
    check(rows, A27, 46, [(9, "as a % of nominal GDP")])
    check(rows, A27, 48, [(8, "Receipts"), (9, "Central govt. income")])
    check(rows, A27, 49, [(9, "as a % of nominal GDP")])
    check(rows, A27, 1, [(9, "Month of financial year end")])
    for r in rows[11:]:
        fy = str(r[0]).strip() if r[0] is not None else ""
        if len(fy) == 7 and fy[4] == "/" and fy[:4].isdigit() and FIRST <= int(fy[:4]) <= LAST:
            y = int(fy[:4])
            assert int(fy[5:]) == (y + 1) % 100 and str(r[1]).startswith("31st March"), (fy, r[1])
            table[y]["fy"] = fy
            table[y]["fy_expenditure_gbp_mn"] = number(r[2])
            table[y]["fy_receipts_gbp_mn"] = number(r[13])
        cy = number(r[43])
        if cy is not None and FIRST <= cy <= LAST:
            table[int(cy)]["cy_expenditure_pct_gdp"] = number(r[46])
            table[int(cy)]["cy_receipts_pct_gdp"] = number(r[49])

    # A29: the national debt, end of calendar year (columns 39-41)
    rows = list(wb[A29].iter_rows(values_only=True))
    check(rows, A29, 39, [(2, "Approximate end calendar year estimates £mn")])
    check(rows, A29, 40, [(4, "Spliced measure of consolidated UK debt including terminable annuities")])
    check(rows, A29, 41, [(4, "as a % of nominal GDP at market prices"),
                          (5, "GDP of GB+all of Ireland up to 1920, GB+N.Ireland from 1921")])
    for r in rows[6:]:
        y = number(r[39])
        if y is not None and FIRST <= y <= LAST:
            table[int(y)]["debt_gbp_mn"] = number(r[40])
            table[int(y)]["debt_pct_gdp"] = number(r[41])

    # A47: consumer prices (column 3, as in data/fetch/boe_millennium_gb27.py)
    rows = list(wb[A47].iter_rows(values_only=True))
    check(rows, A47, 3, [(2, "Consumer Price Index (CPI) -  preferred measure"), (3, "GB/UK(GB+NI)"),
                         (5, "Spliced index 2015=100")])
    for r in rows[6:]:
        y = number(r[0])
        if y is not None and FIRST <= y <= LAST:
            table[int(y)]["cpi_2015_100"] = number(r[3])

    out = []
    for y in years:
        missing = [h for h in HEADER[1:] if table[y].get(h) is None]
        assert not missing, (y, missing)
        out.append([y, table[y]["fy"]] + [round(float(table[y][h]), 6) for h in HEADER[2:]])
    write(OUT, HEADER, out)
    print(f"{len(years)} years {FIRST} - {LAST} -> {OUT}")


if __name__ == "__main__":
    default = Path.home() / ".cache" / "finance-study" / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    p.parent.mkdir(parents=True, exist_ok=True)
    main(p)
