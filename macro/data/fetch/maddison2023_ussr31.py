"""GDP per head of the USSR and six other economies, 1913-1940, from the Maddison Project Database 2023 (Chapter 31).

    python3 data/fetch/maddison2023_ussr31.py          # use (or download) ~/.cache/finance-study/maddison2023
    python3 data/fetch/maddison2023_ussr31.py DIR      # use a copy of mpd2023_web.xlsx already in DIR

The workbook, its URL and its headings are those of data/fetch/maddison2023.py (Chapter 21). The script keeps the
former USSR ('SUN'), the United States, the United Kingdom, Germany and Japan (the industrial economies of the
figure) and Argentina and India (the two economies of Allen's counterfactual, as Davies's review quotes it) for
1913-1940, and writes data/snapshots/maddison2023_ussr.csv (country, year, gdppc_2011usd, pop_thousands; the
population is empty where the database has none: the USSR has no figure for 1914-1919). Every economy has a GDP per
head for every year. The citation policy of the workbook ('Notes' sheet) asks that the original papers be cited when
the data are shown in a graph: the script prints, for each economy, the rows of the sheets 'Sources' and 'Maddison
original sources' whose period covers any of 1913-1940, in the form the snapshot's .json keeps them
('sources_by_country'). For the USSR they are Markevich and Harrison (2011) for 1913-1928 and Moorsteen and Powell
(1966) for 1928-1940, a Western recalculation rather than the official index.
"""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from maddison2023 import HEAD, fetch

OUT = HERE.parent / "snapshots" / "maddison2023_ussr.csv"
FIRST, LAST = 1913, 1940
COUNTRIES = ("SUN", "USA", "GBR", "DEU", "JPN", "ARG", "IND")


def covers(period: str) -> bool:
    """'1928-40', '1913-1928 ', '1874-1940', '1991 -': does the period touch FIRST..LAST?"""
    years = [int(y) for y in re.findall(r"\d{2,4}", period)]
    if not years:
        return False
    a = years[0]
    b = years[1] if len(years) > 1 else 9999
    if b < 100:                                        # '1928-40'
        b += a // 100 * 100
    return a <= LAST and b >= FIRST


def sources(wb) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {c: [] for c in COUNTRIES}
    for sheet in ("Sources", "Maddison original sources"):
        code = None
        for r in wb[sheet].iter_rows(values_only=True):
            if r and isinstance(r[0], str) and re.fullmatch(r"[A-Z]{3}", r[0]):
                code = r[0]
                continue
            if code in out and r and r[0] is None and r[1] and r[2] and covers(str(r[1])):
                out[code].append(f"{sheet}, {str(r[1]).strip()}: {str(r[2]).strip()}")
    return out


def main(dest: Path) -> None:
    wb = openpyxl.load_workbook(fetch(dest), read_only=True, data_only=True)
    rows = wb["Full data"].iter_rows(values_only=True)
    head = tuple(next(rows))[:6]
    if head != HEAD:
        raise ValueError(f"headings {head}, expected {HEAD}")
    got: dict[tuple[str, int], tuple[float, float | None]] = {}
    for r in rows:
        code, _, _, year, gdppc, pop = r[:6]
        if code in COUNTRIES and FIRST <= year <= LAST:
            if gdppc is None:
                raise ValueError(f"{code} {year}: no GDP per head")
            if (code, int(year)) in got:
                raise ValueError(f"{code} {year} appears twice")
            got[(code, int(year))] = (float(gdppc), None if pop is None else float(pop))
    missing = [(c, y) for c in COUNTRIES for y in range(FIRST, LAST + 1) if (c, y) not in got]
    if missing:
        raise ValueError(f"missing {missing[:5]}")
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["country", "year", "gdppc_2011usd", "pop_thousands"])
        for c in COUNTRIES:
            for y in range(FIRST, LAST + 1):
                g, p = got[(c, y)]
                w.writerow([c, y, round(g, 3), "" if p is None else round(p, 3)])
    print(f"{len(got)} rows, {len(COUNTRIES)} economies, {FIRST}-{LAST} -> {OUT}")
    for c, s in sources(wb).items():
        print(c, s)


if __name__ == "__main__":
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / ".cache" / "finance-study" / "maddison2023"
    folder.mkdir(parents=True, exist_ok=True)
    main(folder)
