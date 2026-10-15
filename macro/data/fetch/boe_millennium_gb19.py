"""Great Britain 1790-1830: wholesale and consumer prices and the national debt, from the Bank of England's
"A millennium of macroeconomic data" (Chapter 19).

    python3 data/fetch/boe_millennium_gb19.py           # download the workbook (27 MB) and write the snapshot
    python3 data/fetch/boe_millennium_gb19.py FILE      # use a copy of the workbook already downloaded

Sheet "M6. Mthly prices and wages": the monthly wholesale price index of Gayer, Rostow and Schwartz (1953),
domestic and imported commodities, 1821-5 = 100; the script keeps the years with all twelve months and stores
their mean. Sheet "A47. Wages and prices": the consumer price index, preferred measure (Feinstein 1998 for
1770-1882), spliced, 2015 = 100. Sheet "A29. The National Debt" (par values, £ m, end of financial year): the
total funded and unfunded debt and the spliced total as a per cent of nominal GDP. The script checks the
headings, keeps 1790-1830 and writes data/snapshots/boe_millennium_gb_1790_1830.csv. A financial year
"1815/16" is stored as 1815.
"""
from __future__ import annotations

import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "snapshots" / "boe_millennium_gb_1790_1830.csv"
URL = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/"
       "a-millennium-of-macroeconomic-data-for-the-uk.xlsx")
UA = "Mozilla/5.0 (X11; Linux x86_64) finance-study data fetch"
FIRST, LAST = 1790, 1830
M6, A47, A29 = "M6. Mthly prices and wages", "A47. Wages and prices", "A29. The National Debt"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
# M6: (column, [(row, expected heading)])
GRS = (4, [(4, "Domestic and imported commodities 1790-1850"), (5, "Gayer et al (1953)"), (6, "1821-5=100")])
# annual sheets: (sheet, column, (row, expected heading), name in the snapshot)
COLS = [
    (A47, 3, (2, "Consumer Price Index (CPI) -  preferred measure"), "cpi_2015"),
    (A29, 23, (4, "Total funded and unfunded debt of the UK"), "total_debt_gbp_m"),
    (A29, 43, (4, "as a % of UK nominal GDP according to relevant political boundary"), "debt_pct_gdp"),
]


def year_of(cell) -> int | None:
    if isinstance(cell, (int, float)):
        return int(cell)
    if isinstance(cell, str) and cell[:4].isdigit():
        return int(cell[:4])
    return None


def main(path: Path) -> None:
    if not path.exists():
        req = urllib.request.Request(URL, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=300) as r:
            path.write_bytes(r.read())
    print(f"{path.name}: sha256 {hashlib.sha256(path.read_bytes()).hexdigest()}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list] = {}
    # the monthly wholesale index: the mean of the twelve months
    rows = sheets.setdefault(M6, list(wb[M6].iter_rows(values_only=True)))
    col, heads = GRS
    for row, head in heads:
        if str(rows[row][col]).strip() != head:
            raise ValueError(f"{M6}: column {col} row {row} is {rows[row][col]!r}, expected {head!r}")
    months: dict[int, dict[str, float]] = {}
    for r in rows[8:]:
        y = year_of(r[0])
        if y is not None and FIRST <= y <= LAST and r[1] in MONTHS and isinstance(r[col], (int, float)):
            months.setdefault(y, {})[r[1]] = float(r[col])
    data: dict[str, dict[int, float]] = {
        "wholesale_1821_25": {y: sum(m.values()) / 12 for y, m in months.items() if len(m) == 12}}
    for sheet, col, (row, head), name in COLS:
        rows = sheets.setdefault(sheet, list(wb[sheet].iter_rows(values_only=True)))
        if str(rows[row][col]).strip() != head:
            raise ValueError(f"{sheet}: column {col} row {row} is {rows[row][col]!r}, expected {head!r}")
        s = {}
        for r in rows[row + 1:]:
            y = year_of(r[0])
            if y is not None and FIRST <= y <= LAST and isinstance(r[col], (int, float)):
                s.setdefault(y, float(r[col]))
        data[name] = s
    years = range(FIRST, LAST + 1)
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year"] + list(data))
        for y in years:
            w.writerow([y] + ["" if y not in data[k] else round(data[k][y], 6) for k in data])
    print(f"{len(years)} years -> {OUT}")
    for k, s in data.items():
        print(f"  {k}: {min(s)}-{max(s)}, {len(s)} values")


if __name__ == "__main__":
    default = Path.home() / ".cache" / "finance-study" / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    p.parent.mkdir(parents=True, exist_ok=True)
    main(p)
