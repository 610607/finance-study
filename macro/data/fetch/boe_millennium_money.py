"""Money, output and prices in England, 1490-1700, from the Bank of England's "A millennium of macroeconomic data".

    python3 data/fetch/boe_millennium_money.py           # download the workbook (27 MB) and write the snapshot
    python3 data/fetch/boe_millennium_money.py FILE      # use a copy of the workbook already downloaded

Sheet "A21. GDP per capita 1086+", block "England" (Broadberry, Campbell, Klein, Overton and van Leeuwen 2015):
nominal GDP at market prices (£ m), real GDP at 2013 market prices (£ m), population (m) and the GDP deflator
at market prices (2013 = 100). Sheet "A22. Coin in circulation": total coin in circulation, end of year (£ m),
Palma (2017), "preferred direct method B estimate" (built from mint output, not from GDP, so that velocity
can be computed from it). The script checks the headings, keeps 1490-1700 and writes
data/snapshots/boe_millennium_england_money.csv.
"""
from __future__ import annotations

import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "snapshots" / "boe_millennium_england_money.csv"
URL = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/"
       "a-millennium-of-macroeconomic-data-for-the-uk.xlsx")
UA = "Mozilla/5.0 (X11; Linux x86_64) finance-study data fetch"
FIRST, LAST = 1490, 1700
# (column in sheet A21, heading in row 5, name in the snapshot)
A21 = [(1, "Nominal GDP at market prices", "nominal_gdp_gbp_m"),
       (4, "Real GDP, 2013 market prices", "real_gdp_2013_gbp_m"),
       (6, "Population", "pop_england_m"),
       (9, "GDP deflator at market prices", "gdp_deflator_2013_100")]
A22_COL, A22_HEAD = 2, "Palma (2017), preferred direct method B estimate"


def series(rows, col: int) -> dict[int, float]:
    out = {}
    for r in rows:
        if isinstance(r[0], (int, float)) and FIRST <= r[0] <= LAST and isinstance(r[col], (int, float)):
            out[int(r[0])] = float(r[col])
    return out


def main(path: Path) -> None:
    if not path.exists():
        req = urllib.request.Request(URL, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=300) as r:
            path.write_bytes(r.read())
    print(f"{path.name}: sha256 {hashlib.sha256(path.read_bytes()).hexdigest()}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["A21. GDP per capita 1086+"].iter_rows(values_only=True))
    if rows[3][1] != "England":
        raise ValueError("A21: block 'England' not found")
    for i, title, _ in A21:
        if str(rows[4][i]).strip() != title:
            raise ValueError(f"A21: column {i} is {rows[4][i]!r}, expected {title!r}")
    data = {name: series(rows, i) for i, _, name in A21}
    rows = list(wb["A22. Coin in circulation"].iter_rows(values_only=True))
    if rows[2][A22_COL] != A22_HEAD or rows[3][A22_COL] != "End year":
        raise ValueError(f"A22: column {A22_COL} is {rows[2][A22_COL]!r}")
    data["coin_palma_gbp_m"] = series(rows, A22_COL)
    years = range(FIRST, LAST + 1)
    missing = [(k, y) for k, s in data.items() for y in years if y not in s]
    if missing:
        raise ValueError(f"missing values: {missing[:10]}")
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year"] + list(data))
        for y in years:
            w.writerow([y] + [round(data[k][y], 6) for k in data])
    print(f"{len(years)} years -> {OUT}")


if __name__ == "__main__":
    default = Path.home() / ".cache" / "finance-study" / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    p.parent.mkdir(parents=True, exist_ok=True)
    main(p)
