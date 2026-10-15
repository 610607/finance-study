"""England 1270-1913: population and real GDP per head, from the Bank of England's "A millennium of
macroeconomic data" (Chapter 21).

    python3 data/fetch/boe_millennium_gb21.py           # download the workbook (27 MB) and write the snapshot
    python3 data/fetch/boe_millennium_gb21.py FILE      # use a copy of the workbook already downloaded

Sheet "A2. Pop of Eng & GB 1086-1870" (Broadberry, Campbell, Klein, Overton and van Leeuwen 2015): population of
England, millions (to 1870). Sheet "A21. GDP per capita 1086+", the block headed "England": real GDP per head at
2013 market prices, £ (Broadberry et al. 2015 to 1700, spliced by the Bank to the later series for England). The
script checks the headings (and the country heading of the block), keeps 1270-1913 and writes
data/snapshots/boe_millennium_england_1270_1913.csv.
"""
from __future__ import annotations

import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "snapshots" / "boe_millennium_england_1270_1913.csv"
URL = ("https://www.bankofengland.co.uk/-/media/boe/files/statistics/research-datasets/"
       "a-millennium-of-macroeconomic-data-for-the-uk.xlsx")
UA = "Mozilla/5.0 (X11; Linux x86_64) finance-study data fetch"
FIRST, LAST = 1270, 1913
A2, A21 = "A2. Pop of Eng & GB 1086-1870", "A21. GDP per capita 1086+"
# (sheet, column, (row, expected heading), name in the snapshot)
COLS = [
    (A2, 1, (7, "Population of England, millions"), "pop_england_m"),
    (A21, 8, (4, "Real GDP per capita"), "england_real_gdp_pc_2013gbp"),
]
BLOCKS = [(A21, 3, 1, "England")]


def year_of(cell) -> int | None:
    if isinstance(cell, (int, float)):
        return int(cell)
    if isinstance(cell, str) and cell[:4].isdigit():
        return int(cell[:4])
    return None


def main(path: Path) -> None:
    if not path.exists():
        req = urllib.request.Request(URL, headers={"User-Agent": UA})
        with urllib.request.urlopen(req, timeout=300) as r:
            path.write_bytes(r.read())
    print(f"{path.name}: sha256 {hashlib.sha256(path.read_bytes()).hexdigest()}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list] = {}
    for sheet, row, col, head in BLOCKS:
        rows = sheets.setdefault(sheet, list(wb[sheet].iter_rows(values_only=True)))
        if not str(rows[row][col]).startswith(head):
            raise ValueError(f"{sheet}: column {col} row {row} is {rows[row][col]!r}, expected {head!r}")
    data: dict[str, dict[int, float]] = {}
    for sheet, col, (row, head), name in COLS:
        rows = sheets.setdefault(sheet, list(wb[sheet].iter_rows(values_only=True)))
        if str(rows[row][col]).strip() != head:
            raise ValueError(f"{sheet}: column {col} row {row} is {rows[row][col]!r}, expected {head!r}")
        s = {}
        for r in rows[row + 1:]:
            y = year_of(r[0])
            if y is not None and FIRST <= y <= LAST and isinstance(r[col], (int, float)):
                s.setdefault(y, float(r[col]))
        data[name] = s
    years = range(FIRST, LAST + 1)
    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year"] + list(data))
        for y in years:
            w.writerow([y] + ["" if y not in data[k] else round(data[k][y], 6) for k in data])
    print(f"{len(years)} years -> {OUT}")
    for k, s in data.items():
        print(f"  {k}: {min(s)}-{max(s)}, {len(s)} values")


if __name__ == "__main__":
    default = Path.home() / ".cache" / "finance-study" / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else default
    p.parent.mkdir(parents=True, exist_ok=True)
    main(p)
